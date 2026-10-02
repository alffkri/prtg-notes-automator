import os
import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Inches

# Konfigurasi Halaman Streamlit
st.set_page_config(
    page_title="PRTG Word Notes Automator", 
    page_icon="📝", 
    layout="centered",
    initial_sidebar_state="collapsed"
)

# Custom CSS untuk merapikan layout, card, dan styling responsive
st.markdown("""
    <style>
    .main {
        background-color: #0e1117;
        color: #fafafa;
    }
    .stButton>button {
        width: 100%;
        background-color: #ff4b4b;
        color: white;
        font-weight: bold;
        border-radius: 8px;
        padding: 0.6rem 1rem;
        border: none;
        transition: all 0.3s ease;
    }
    .stButton>button:hover {
        background-color: #ff2b2b;
        box-shadow: 0 4px 12px rgba(255, 75, 75, 0.4);
    }
    h1 {
        font-family: 'Inter', sans-serif;
        font-weight: 800;
        font-size: 2.2rem;
        background: linear-gradient(90deg, #ff4b4b, #ffa15c);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        text-align: center;
        margin-bottom: 0px;
    }
    .subtitle {
        text-align: center;
        color: #8b949e;
        font-size: 0.95rem;
        margin-bottom: 2rem;
    }
    .card-container {
        background-color: #161b22;
        border: 1px solid #30363d;
        border-radius: 12px;
        padding: 20px;
        height: 100%;
        box-sizing: border-box;
    }
    .footer {
        text-align: center;
        color: #8b949e;
        font-size: 0.85rem;
        margin-top: 3rem;
        padding-top: 1rem;
        border-top: 1px solid #30363d;
    }
    </style>
""", unsafe_allow_html=True)

# Header Utama
st.markdown("<h1>📝 PRTG Excel-to-Word Notes Automator</h1>", unsafe_allow_html=True)
st.markdown("<p class='subtitle'>Aplikasi web otomatis profesional untuk menyinkronkan data Downtime, Uptime, dan Keterangan dari Excel ke bagian <b>Note :</b> dokumen Word secara presisi.</p>", unsafe_allow_html=True)

# Layout dengan Kolom dan Container Rapi
col1, col2 = st.columns(2, gap="medium")

with col1:
    st.markdown("""
        <div class="card-container">
            <h3 style="margin-top:0; font-size: 1.05rem; color: #ffffff;">📄 1. Dokumen Word Mentah</h3>
            <p style="font-size: 0.8rem; color: #8b949e; margin-bottom: 10px;">Pilih file Word hasil generate PRTG (.docx)</p>
        """, unsafe_allow_html=True)
    word_file = st.file_uploader("Upload Word", type=["docx"], key="word", label_visibility="collapsed")
    if word_file:
        st.success(f"Terpilih: **{word_file.name}**")
    st.markdown("</div>", unsafe_allow_html=True)

with col2:
    st.markdown("""
        <div class="card-container">
            <h3 style="margin-top:0; font-size: 1.05rem; color: #ffffff;">📊 2. File Excel Olahan</h3>
            <p style="font-size: 0.8rem; color: #8b949e; margin-bottom: 10px;">Pilih file Excel data catatan (.xlsx)</p>
        """, unsafe_allow_html=True)
    excel_file = st.file_uploader("Upload Excel", type=["xlsx", "xls"], key="excel", label_visibility="collapsed")
    if excel_file:
        st.success(f"Terpilih: **{excel_file.name}**")
    st.markdown("</div>", unsafe_allow_html=True)

st.markdown("<br>", unsafe_allow_html=True)

# Fungsi inti penggabungan Note dengan Remove Space Before & After pada Paragraf Note
def process_word_notes(word_path, excel_path, output_path):
    try:
        excel_file_obj = pd.ExcelFile(excel_path)
        doc = Document(word_path)
    except Exception as e:
        return False, str(e)

    data_notes = {}
    for sheet_name in excel_file_obj.sheet_names:
        df_sheet = pd.read_excel(excel_file_obj, sheet_name=sheet_name)
        
        required_cols = ['LOKASI', 'Downtime', 'Uptime', 'Keterangan']
        if not all(col in df_sheet.columns for col in required_cols):
            continue
            
        data_notes[sheet_name] = {}
        current_lokasi = None
        
        for _, row in df_sheet.iterrows():
            lokasi_val = str(row['LOKASI']).strip()
            if lokasi_val and lokasi_val.lower() != "nan" and lokasi_val != "-":
                current_lokasi = lokasi_val
                
            if not current_lokasi:
                continue
                
            downtime = str(row['Downtime']).strip()
            uptime = str(row['Uptime']).strip()
            keterangan = str(row['Keterangan']).strip()
            
            if downtime and downtime != "-" and downtime.lower() != "nan":
                detail_teks = f"{downtime}     {uptime}"
                if keterangan and keterangan.lower() != "nan" and keterangan != "-":
                    detail_teks += f"     {keterangan}"
                
                if current_lokasi not in data_notes[sheet_name]:
                    data_notes[sheet_name][current_lokasi] = []
                data_notes[sheet_name][current_lokasi].append(detail_teks)
            elif keterangan and keterangan.lower() != "nan" and keterangan != "-":
                if current_lokasi not in data_notes[sheet_name]:
                    data_notes[sheet_name][current_lokasi] = []
                data_notes[sheet_name][current_lokasi].append(keterangan)

    current_sheet = None
    current_location = None

    for element in doc.element.body:
        if element.tag.endswith('p'):
            from docx.text.paragraph import Paragraph
            p = Paragraph(element, doc)
            text_p = p.text.strip()
            
            if p.style.name.startswith('Heading 1'):
                current_sheet = text_p.split(".", 1)[1].strip() if "." in text_p else text_p
                current_location = None
            elif p.style.name.startswith('Heading 2'):
                current_location = text_p
                
        elif element.tag.endswith('tbl'):
            from docx.table import Table
            table = Table(element, doc)
            table_text = "".join([cell.text for row in table.rows for cell in row.cells])
            
            if "Note" in table_text:
                matched_notes = []
                if current_sheet in data_notes:
                    for loc_key, notes_list in data_notes[current_sheet].items():
                        if current_location and (loc_key.lower() in current_location.lower() or current_location.lower() in loc_key.lower()):
                            matched_notes = notes_list
                            break
                
                if matched_notes:
                    for row in table.rows:
                        for cell in row.cells:
                            if "Note" in cell.text:
                                paragraphs_to_remove = []
                                found_note_section = False
                                for p in cell.paragraphs:
                                    if "Note" in p.text:
                                        found_note_section = True
                                    if found_note_section:
                                        paragraphs_to_remove.append(p)
                                        
                                for p in paragraphs_to_remove:
                                    p_elem = p._p
                                    p_elem.getparent().remove(p_elem)
                                
                                p_note_title = cell.add_paragraph()
                                # REMOVE SPACE BEFORE & AFTER untuk judul Note
                                p_note_title.paragraph_format.space_before = Pt(0)
                                p_note_title.paragraph_format.space_after = Pt(0)
                                run_title = p_note_title.add_run("Note :")
                                run_title.font.name = 'Times New Roman'
                                run_title.font.size = Pt(12)
                                run_title.font.bold = False
                                run_title.font.color.rgb = RGBColor(0, 0, 0)
                                
                                for note_text in matched_notes:
                                    p_bullet = cell.add_paragraph()
                                    p_bullet.paragraph_format.left_indent = Inches(0.35)
                                    p_bullet.paragraph_format.first_line_indent = Inches(-0.2)
                                    # REMOVE SPACE BEFORE & AFTER untuk baris-baris note agar rapat
                                    p_bullet.paragraph_format.space_before = Pt(0)
                                    p_bullet.paragraph_format.space_after = Pt(0)
                                    
                                    run_bullet = p_bullet.add_run(f"-  {note_text}")
                                    run_bullet.font.name = 'Times New Roman'
                                    run_bullet.font.size = Pt(12)
                                    run_bullet.font.color.rgb = RGBColor(0, 0, 0)

    doc.save(output_path)
    return True, "Berhasil"

# Tombol Eksekusi Utama
if st.button("🚀 Proses Penggabungan Note Sekarang", type="primary"):
    if word_file is not None and excel_file is not None:
        with st.spinner("⏳ Sedang memproses dokumen Word dan mencocokkan data Excel secara presisi..."):
            temp_word_path = "temp_input.docx"
            temp_excel_path = "temp_input.xlsx"
            output_word_path = "Report_PRTG_Final_Dengan_Note.docx"
            
            with open(temp_word_path, "wb") as f:
                f.write(word_file.getbuffer())
            with open(temp_excel_path, "wb") as f:
                f.write(excel_file.getbuffer())
                
            success, msg = process_word_notes(temp_word_path, temp_excel_path, output_word_path)
            
            if success:
                st.success("🎉 Penggabungan berhasil dilakukan dengan sempurna!")
                st.markdown("<br>", unsafe_allow_html=True)
                
                with open(output_word_path, "rb") as f:
                    st.download_button(
                        label="📥 Download Laporan Word Final (.docx)",
                        data=f,
                        file_name="Report_PRTG_Final_Dengan_Note.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
            else:
                st.error(f"❌ Terjadi kesalahan saat memproses: {msg}")
    else:
        st.warning("⚠️ Mohon unggah file Word mentah dan file Excel olahan terlebih dahulu pada kolom di atas!")

# Footer Copyright
st.markdown("""
    <div class="footer">
        © 2026 M Alif Fikri. All rights reserved.
    </div>
""", unsafe_allow_html=True)